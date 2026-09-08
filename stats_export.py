import html
import os


PDF_FONT_NAME = "KoreanStatsFont"


def export_excel(file_path, title, details, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    for index, (sheet_name, headers, rows) in enumerate(tables):
        worksheet = workbook.worksheets[0] if index == 0 else workbook.create_sheet()
        worksheet.title = sheet_name[:31]
        worksheet.append([title])
        worksheet.merge_cells(
            start_row=1, start_column=1, end_row=1, end_column=len(headers)
        )
        title_cell = worksheet.cell(1, 1)
        title_cell.font = Font(size=16, bold=True, color="1F4E78")

        row_number = 3
        for label, value in details:
            worksheet.cell(row_number, 1, label)
            worksheet.cell(row_number, 2, value)
            worksheet.cell(row_number, 1).font = Font(bold=True)
            row_number += 1

        header_row = row_number + 1
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(header_row, column, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4C72B0")
            cell.alignment = Alignment(horizontal="center")

        for row in rows:
            worksheet.append(list(row))

        last_row = header_row + len(rows)
        worksheet.freeze_panes = f"A{header_row + 1}"
        worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"

        for column, header in enumerate(headers, start=1):
            values = [header] + [row[column - 1] for row in rows]
            width = max(len(str(value or "")) for value in values) + 3
            worksheet.column_dimensions[get_column_letter(column)].width = min(
                max(width, 10), 40
            )
            if rows:
                for cell in worksheet.iter_cols(
                    min_col=column,
                    max_col=column,
                    min_row=header_row + 1,
                    max_row=last_row,
                ):
                    for item in cell:
                        item.alignment = Alignment(
                            horizontal="center" if column > 1 else "left"
                        )

    workbook.save(file_path)


def export_pdf(file_path, title, details, tables):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Flowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font_path = _find_korean_font()
    if PDF_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, font_path))

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "KoreanTitle",
        parent=styles["Title"],
        fontName=PDF_FONT_NAME,
        fontSize=18,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1F4E78"),
    )
    heading_style = ParagraphStyle(
        "KoreanHeading",
        parent=styles["Heading2"],
        fontName=PDF_FONT_NAME,
        fontSize=12,
        leading=16,
        spaceBefore=8,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "KoreanBody",
        parent=styles["BodyText"],
        fontName=PDF_FONT_NAME,
        fontSize=9,
        leading=14,
    )

    document = SimpleDocTemplate(
        file_path,
        pagesize=landscape(A4),
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=title,
    )
    story: list[Flowable] = [
        Paragraph(html.escape(title), title_style),
        Spacer(1, 5 * mm),
    ]
    detail_text = "<br/>".join(
        f"<b>{html.escape(str(label))}</b>: {html.escape(str(value))}"
        for label, value in details
    )
    if detail_text:
        story.extend([Paragraph(detail_text, body_style), Spacer(1, 4 * mm)])

    for section_name, headers, rows in tables:
        story.append(Paragraph(html.escape(section_name), heading_style))
        data = [list(headers)] + [list(row) for row in rows]
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), PDF_FONT_NAME),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4C72B0")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C9D6")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EDF3F7")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.extend([table, Spacer(1, 5 * mm)])

    document.build(story)


def _find_korean_font():
    windir = os.environ.get("WINDIR", "C:\\Windows")
    candidates = [
        os.path.join(windir, "Fonts", "malgun.ttf"),
        "/mnt/c/Windows/Fonts/malgun.ttf",
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    ]
    for path in candidates:
        if os.path.isfile(path):
            return path
    raise RuntimeError(
        "PDF 한글 출력을 위한 맑은 고딕 또는 Noto Sans CJK 글꼴을 찾을 수 없습니다."
    )
